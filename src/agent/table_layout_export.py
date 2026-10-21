"""Export the exact validated layout; no provider calls or executable formulas."""
import csv
import io
import textwrap
import re

def safe(value):
    if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@')):
        if re.fullmatch(r'[+-]?(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?',value):return value
        return "'"+value
    return value

def footnotes(doc):
    return [*[f"주 {n['number']}: {n['text']}" for n in doc.get('annotations',[])],*doc['notes'],
            *[f"자료: {s.get('displaySource','KOSIS')} · {s['tableName']}" for s in doc['sources']]]

def render_layout(doc,kind):
    if kind=='csv':
        out=io.StringIO();writer=csv.writer(out)
        writer.writerow([safe(doc['caption'])])
        for s in doc['layoutSections']:
            if s.get('sheetRows'):
                for r in s['sheetRows']:
                    row=['']*len(s['columns'])
                    for c in r['cells']:row[c['column']]=safe(c['text'])
                    writer.writerow(row)
                continue
            for headers in s['headerRows']:
                row=['']*len(s['columns'])
                for c in headers:row[c['column']]=safe(c['text'])
                writer.writerow(row)
            for r in s['layoutRows']:writer.writerow([safe(r['label'])] if r.get('group') else [safe(r['display'][c['id']]) for c in s['columns']])
        for line in footnotes(doc):writer.writerow([safe(line)])
        return out.getvalue().encode('utf-8-sig')
    if kind=='xlsx':return xlsx(doc)
    if kind not in {'pdf','png'}:raise ValueError('지원하지 않는 표 파일 형식입니다.')
    return image(doc,kind)

def xlsx(doc):
    from openpyxl import Workbook
    from openpyxl.styles import Font,Alignment,Border,Side,PatternFill
    from openpyxl.utils import get_column_letter
    book=Workbook();book.remove(book.active);evidence=book.create_sheet('계산근거')
    evidence.append(['표','행','열','원자료 계열','주기','시점','단위','값','기준 A','현재 B','연산','보조 증가율 (%)','중단 사유'])
    def put(sheet,row,col,value):
        cell=sheet.cell(row,col,value)
        if isinstance(value,str):cell.data_type='s'
        return cell
    def style(cell,f):
        cell.font=Font(name='맑은 고딕',size=f.get('font_size',11),bold=f.get('bold',False),italic=f.get('italic',False),underline='single' if f.get('underline') else None,color=f.get('color','#17324d')[1:])
        cell.fill=PatternFill('solid',fgColor=f.get('background','#ffffff')[1:])
        cell.alignment=Alignment(horizontal=f.get('align','right'),vertical={'middle':'center'}.get(f.get('vertical_align','middle'),f.get('vertical_align','middle')),wrap_text=f.get('wrap',True))
        from table_borders import excel_border
        sides={k:Side(style=excel_border(f.get(k,'none')),color=f.get(k+'_color','#577088')[1:]) for k in ['top','bottom','left','right']}
        cell.border=Border(**sides)
        decimals=f.get('decimals');decimals=doc['options']['decimals'] if decimals is None else decimals
        number='#,##0'+('.'+'0'*decimals if decimals else '')
        cell.number_format=('+'+number+';-'+number+';'+number) if f.get('signed') else number
        if f.get('zero_dash'):cell.number_format=cell.number_format.rsplit(';',1)[0]+';"−"' if f.get('signed') else number+';-'+number+';"−"'
    for i,s in enumerate(doc['layoutSections']):
        ws=book.create_sheet(f'표{i+1}',i);n=len(s['columns'])
        put(ws,1,1,doc['caption']);ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=n)
        ws.row_dimensions[1].height=30;ws.cell(1,1).font=Font(name='맑은 고딕',bold=True,size=15)
        for ci,c in enumerate(s['columns'],1):ws.column_dimensions[get_column_letter(ci)].width=c['width']/7
        for hi,headers in enumerate(s['headerRows'],3):
            ws.row_dimensions[hi].height=32
            for h in headers:
                ci=h['column']+1;cell=put(ws,hi,ci,h['text']);style(cell,h.get('format',{'bold':True,'align':'center','background':'#eaf2f8'}))
                if h['rowSpan']>1 or h['colSpan']>1:ws.merge_cells(start_row=hi,start_column=ci,end_row=hi+h['rowSpan']-1,end_column=ci+h['colSpan']-1)
        start=3+len(s['headerRows']);ws.freeze_panes=f'B{start}'
        if doc['options']['layout']['repeat_headers']:ws.print_title_rows=f'1:{start-1}'
        for ri,r in enumerate(s['layoutRows'],start):
            ws.row_dimensions[ri].height=r.get('height',32)*.75
            if r.get('group'):
                put(ws,ri,1,r['label']);ws.merge_cells(start_row=ri,start_column=1,end_row=ri,end_column=n);style(ws.cell(ri,1),{'bold':True,'align':'left','background':'#edf3f7','top':'thin'});continue
            for ci,c in enumerate(s['columns'],1):
                cid=c['id'];info=r.get('comparisons',{}).get(cid);value=r['label'] if cid=='label' else r['cells'].get(cid)
                if info and info['operation']=='difference_growth':value=r['display'][cid]
                cell=put(ws,ri,ci,value);style(cell,r['formats'][cid]);cell.alignment=Alignment(horizontal=r['formats'][cid]['align'],vertical='center',wrap_text=True,indent=r.get('level',0) if cid=='label' else 0)
                markers=r['markers'].get(cid,[])
                if markers:
                    from openpyxl.comments import Comment
                    cell.comment=Comment('\n'.join(footnotes(doc)[m-1] for m in markers),'StatBridge')
                if cid!='label':
                    ref=r.get('references',{}).get(cid,{})
                    data=[ws.title,r['label'],c['label'],ref.get('series',r.get('sourceLabel','')),ref.get('frequency',''),ref.get('period',''),ref.get('unit',info.get('unit','') if info else ''),r['cells'].get(cid),info.get('observations',[None,None])[0] if info else None,info.get('observations',[None,None])[1] if info else None,info.get('operation','') if info else '',info.get('secondary') if info else None,info.get('reason','') if info else '']
                    for j,v in enumerate(data,1):put(evidence,evidence.max_row+1 if j==1 else er,j,v);er=evidence.max_row
        end=start+len(s['layoutRows'])
        if s.get('sheetRows'):
            for merged in list(ws.merged_cells.ranges):
                if merged.min_row>=3 and merged.max_row<end:ws.unmerge_cells(str(merged))
            for row in ws.iter_rows(min_row=3,max_row=end-1,max_col=n):
                for cell in row:cell.value=None
            source_rows={r['id']:r for r in s['layoutRows']}
            from openpyxl.comments import Comment
            for ri,row in enumerate(s['sheetRows'],3):
                ws.row_dimensions[ri].height=row['height']*.75
                for c in row['cells']:
                    ci=c['column']+1;value=c['text'];raw=source_rows.get(row['id'])
                    if raw and not raw.get('group') and c['rowSpan']==c['colSpan']==1 and ci>1:
                        cid=s['columns'][ci-1]['id'];info=raw.get('comparisons',{}).get(cid)
                        if not (info and info['operation']=='difference_growth'):value=raw['cells'].get(cid)
                    cell=put(ws,ri,ci,value);style(cell,c['format'])
                    if c['rowSpan']>1 or c['colSpan']>1:
                        ws.merge_cells(start_row=ri,start_column=ci,end_row=ri+c['rowSpan']-1,end_column=ci+c['colSpan']-1)
                        cell.comment=Comment('표시 셀 병합입니다. 원관측값은 계산근거 시트에 모두 보존되어 있으며 합산하지 않았습니다.','StatBridge')
        for j,line in enumerate(footnotes(doc),end+1):put(ws,j,1,line);ws.merge_cells(start_row=j,start_column=1,end_row=j,end_column=n);ws.cell(j,1).alignment=Alignment(wrap_text=True,vertical='top');ws.row_dimensions[j].height=max(25,len(line)//60*15)
        ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='portrait' if doc['options']['print_preset']=='a4_portrait' else 'landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.print_options.horizontalCentered=True;ws.print_area=ws.dimensions
        from openpyxl.worksheet.pagebreak import Break
        count=doc['options']['layout']['rows_per_page']
        for j in range(start+count,end,count):ws.row_breaks.append(Break(id=j-1))
    out=io.BytesIO();book.save(out);return out.getvalue()

def image(doc,kind):
    if all(s.get('sheetRows') for s in doc['layoutSections']):
        from table_sheet_image import render_sheet_image
        return render_sheet_image(doc,kind)
    from report_export import EXPORT_LOCK,_font
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages
    from matplotlib.patches import Rectangle
    from PIL import Image
    pages=[];preset=doc['options']['print_preset'];limit=min(doc['options']['layout']['rows_per_page'],12 if preset=='a4_landscape' else 20 if preset=='a4_portrait' else 60)
    capacity=720 if preset=='a4_portrait' else 1000 if preset=='a4_landscape' else 1400
    for original in doc['layoutSections']:
        # Repeat the identifier column in horizontal panels. Project spanning
        # headers into each panel instead of clipping or losing their hierarchy.
        panels=[];indices=[0];used=original['columns'][0]['width']
        for ci,c in enumerate(original['columns'][1:],1):
            if len(indices)>1 and used+c['width']>capacity:panels.append(indices);indices=[0];used=original['columns'][0]['width']
            indices.append(ci);used+=c['width']
        if indices:panels.append(indices)
        for indices in panels:
            s={**original,'columns':[original['columns'][i] for i in indices]};headers=[]
            for row in original['headerRows']:
                cells=[]
                for h in row:
                    hit=[i for i in indices if h['column']<=i<h['column']+h['colSpan']]
                    if hit:cells.append({**h,'column':indices.index(hit[0]),'colSpan':len(hit)})
                headers.append(cells)
            s['headerRows']=headers;chunk=[];first=True
            for r in s['layoutRows']:
                if len(chunk)>=limit and not chunk[-1].get('group'):
                    pages.append(({**s,'headerRows':headers if first or doc['options']['layout']['repeat_headers'] else []},chunk));chunk=[];first=False
                chunk.append(r)
            if chunk:pages.append(({**s,'headerRows':headers if first or doc['options']['layout']['repeat_headers'] else []},chunk))
    if not pages or len(pages)>120:raise ValueError('출력할 표가 없거나 120쪽을 초과합니다.')
    out=io.BytesIO();images=[];pixels=0
    with EXPORT_LOCK:
        font=_font();pdf=PdfPages(out) if kind=='pdf' else None
        try:
            for page,(s,rows) in enumerate(pages):
                widths=[c['width'] for c in s['columns']];total=sum(widths)
                if total>2400:raise ValueError('인쇄 너비가 너무 큽니다. 열을 줄이거나 Excel로 내보내 주세요.')
                xs=[0]
                for w in widths:xs.append(xs[-1]+w)
                hhead=len(s['headerRows'])*42;heights=[r.get('height',34) for r in rows];foot=[line for text in footnotes(doc) for line in textwrap.wrap(text,100)]
                height=110+hhead+sum(heights)+len(foot)*17+30
                if height>3000:raise ValueError('인쇄 높이가 너무 큽니다. 페이지당 행 수를 줄여 주세요.')
                width=max(850,total+60);figsize=(8.27,11.69) if preset=='a4_portrait' else (11.69,8.27) if preset=='a4_landscape' else (width/100,height/100)
                fig=plt.figure(figsize=figsize);ax=fig.add_axes([.03,.02,.94,.96]);ax.set_xlim(0,total);ax.set_ylim(height,0);ax.axis('off')
                ax.text(0,15,doc['caption'].replace('−','-'),fontproperties=font,fontsize=15,va='top');ax.text(0,48,f"{s['periodCaption']} · {page+1}/{len(pages)}",fontproperties=font,fontsize=9,va='top')
                fig.canvas.draw();renderer=fig.canvas.get_renderer()
                def cell(x,y,w,h,text,fmt):
                    ax.add_patch(Rectangle((x,y),w,h,facecolor=fmt.get('background','#ffffff'),edgecolor='#cbd8e4',linewidth=.35))
                    for edge,coords in {'top':([x,x+w],[y,y]),'bottom':([x,x+w],[y+h,y+h]),'left':([x,x],[y,y+h]),'right':([x+w,x+w],[y,y+h])}.items():
                        style=fmt.get(edge,'none')
                        if style!='none':ax.plot(*coords,color=fmt.get('color','#17324d'),linewidth=1.4 if style=='medium' else .7,linestyle=':' if style=='dotted' else '-')
                    align=fmt.get('align','right');original=str(text).replace('−','-');size=fmt.get('font_size',11)
                    available=abs(ax.transData.transform((w-12,0))[0]-ax.transData.transform((0,0))[0]);vertical=abs(ax.transData.transform((0,h-8))[1]-ax.transData.transform((0,0))[1])
                    def wrap(value,prop):
                        lines=[];line=''
                        for char in value:
                            if char=='\n' or (line and renderer.get_text_width_height_descent(line+char,prop,False)[0]>available):lines.append(line);line=''
                            if char!='\n':line+=char
                        return '\n'.join([*lines,line])
                    while True:
                        prop=font.copy();prop.set_size(size);prop.set_weight('bold' if fmt.get('bold') else 'normal')
                        text=wrap(original,prop) if fmt.get('wrap',True) else original
                        width=max(renderer.get_text_width_height_descent(line,prop,False)[0] for line in text.split('\n'))
                        if (width<=available and len(text.split('\n'))*size*fig.dpi/72*1.2<=vertical) or size<=6:break
                        size-=.5
                    if width>available or len(text.split('\n'))*size*fig.dpi/72*1.2>vertical:raise ValueError('셀 내용이 인쇄 영역을 넘습니다. 열 너비·행 높이를 늘리거나 표시 이름을 줄여 주세요.')
                    ax.text(x+(6 if align=='left' else w/2 if align=='center' else w-6),y+h/2,text,fontproperties=prop,color=fmt.get('color','#17324d'),ha=align,va='center',clip_on=True)
                for hi,headers in enumerate(s['headerRows']):
                    for h in headers:cell(xs[h['column']],82+hi*42,sum(widths[h['column']:h['column']+h['colSpan']]),42*h['rowSpan'],h['text']+(' '+','.join(str(n)+')' for n in h.get('markers',[])) if h.get('markers') else ''),h.get('format',{'bold':True,'align':'center','background':'#eaf2f8'}))
                y=82+hhead
                for r,h in zip(rows,heights):
                    if r.get('group'):cell(0,y,total,h,'  '*r['level']+r['label'],{'bold':True,'align':'left','background':'#edf3f7','top':'thin'})
                    else:
                        for ci,c in enumerate(s['columns']):
                            cid=c['id'];label=r['display'][cid]+(' '+','.join(str(n)+')' for n in r['markers'].get(cid,[])) if r['markers'].get(cid) else '')
                            cell(xs[ci],y,widths[ci],h,('  '*r['level'] if cid=='label' else '')+label,r['formats'][cid])
                    y+=h
                for j,line in enumerate(foot):ax.text(0,y+12+j*17,line.replace('−','-'),fontproperties=font,fontsize=8,va='top')
                try:
                    if pdf:pdf.savefig(fig)
                    else:
                        buf=io.BytesIO();fig.savefig(buf,format='png',dpi=120);buf.seek(0);im=Image.open(buf).convert('RGB');pixels+=im.width*im.height
                        if pixels>40_000_000:
                            im.close();raise ValueError('이미지가 너무 큽니다. PDF로 내보내 주세요.')
                        images.append(im)
                finally:plt.close(fig)
            if pdf:pdf.close();return out.getvalue()
            if sum(im.width*im.height for im in images)>40_000_000:raise ValueError('이미지가 너무 큽니다. PDF로 내보내 주세요.')
            combined=Image.new('RGB',(max(im.width for im in images),sum(im.height for im in images)), 'white');y=0
            for im in images:combined.paste(im,(0,y));y+=im.height;im.close()
            combined.save(out,format='PNG');combined.close();return out.getvalue()
        finally:
            if pdf:pdf.close()
            for im in images:im.close()
