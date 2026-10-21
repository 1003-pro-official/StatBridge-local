"""Bounded, non-executable tabular file decoding into one shared input contract."""
from statbridge_mcp.observability import observed
import base64, csv, io, json, zipfile, zlib, re
from pathlib import Path
from expression_tree_validator import DerivedValidationError

FIELDS = ['period', 'value', 'unit', 'frequency', 'source']
MAX_BYTES = 2_000_000
FORMATS = ['csv', 'json', 'tsv', 'xlsx']

def unique_object(pairs):
    result = {}
    for k,v in pairs:
        if k in result: raise DerivedValidationError('JSON의 중복 필드 이름은 허용하지 않습니다.')
        result[k] = v
    return result

@observed('file.parse', 'span')
def decode_file(filename, content, encoding='utf8'):
    suffix = Path(filename).suffix.lower()[1:]
    if suffix not in FORMATS or len(filename)>120 or any(c in filename for c in '/\\\x00'):
        raise DerivedValidationError('CSV, JSON, TSV 또는 XLSX 파일을 선택해 주세요.')
    if encoding not in ('utf8','base64'):
        raise DerivedValidationError('파일 인코딩이 유효하지 않습니다.')
    if len(content)>2_700_000: raise DerivedValidationError('파일은 2MB 이하여야 합니다.')
    try: raw = base64.b64decode(content, validate=True) if encoding=='base64' else content.encode('utf8')
    except (ValueError, UnicodeError) as exc: raise DerivedValidationError('파일 내용을 읽을 수 없습니다.') from exc
    if len(raw)>MAX_BYTES: raise DerivedValidationError('파일은 2MB 이하여야 합니다.')
    try:
        if suffix=='xlsx':
            if encoding!='base64': raise DerivedValidationError('XLSX 파일은 바이너리 업로드를 사용해 주세요.')
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                entries=archive.infolist()
                if len(entries)>128 or sum(e.file_size for e in entries)>8_000_000 or any(e.flag_bits & 1 for e in entries):
                    raise DerivedValidationError('엑셀 파일의 압축 해제 용량 한도를 초과했거나 암호화된 파일입니다.')
                if any('vba' in e.filename.lower() or 'externallinks' in e.filename.lower() for e in entries):
                    raise DerivedValidationError('매크로와 외부 링크가 포함된 엑셀 파일은 사용할 수 없습니다.')
                for e in entries:
                    if e.filename.endswith(('.xml','.rels')):
                        xml=archive.read(e).replace(b'\x00',b'').upper()
                        if b'<!DOCTYPE' in xml or b'<!ENTITY' in xml or re.search(rb'''TARGETMODE\s*=\s*['"]EXTERNAL['"]''',xml):
                            raise DerivedValidationError('외부 참조 또는 엔터티가 포함된 엑셀 파일은 사용할 수 없습니다.')
            from openpyxl import load_workbook
            book=load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
            try:
                if len(book.worksheets)!=1: raise DerivedValidationError('엑셀은 데이터 시트 하나만 포함해 주세요.')
                sheet=book.worksheets[0]
                if not sheet.max_row or not sheet.max_column:
                    raise DerivedValidationError('행·열 범위가 명시된 XLSX 파일로 저장해 주세요.')
                if sheet.max_row>6001 or sheet.max_column>5:
                    raise DerivedValidationError('엑셀은 필수 열 5개와 최대 6,000행만 사용할 수 있습니다.')
                values=[]
                for cells in sheet.iter_rows():
                    if len(values)>6000 or len(cells)!=5: raise DerivedValidationError('엑셀 행 또는 열 한도를 초과했습니다.')
                    if any(c.data_type in ('f','e') for c in cells): raise DerivedValidationError('엑셀 수식·오류 셀은 값으로 저장한 뒤 올려 주세요.')
                    values.append([c.value for c in cells])
                if not values or values.pop(0)!=FIELDS: raise DerivedValidationError('엑셀 첫 행은 period,value,unit,frequency,source 순서여야 합니다.')
                rows=[dict(zip(FIELDS, r)) for r in values]
            finally: book.close()
        else:
            text=raw.decode('utf-8-sig')
            if '\ufffd' in text: raise DerivedValidationError('UTF-8로 저장한 파일을 사용해 주세요.')
            if suffix=='json':
                rows=json.loads(text,object_pairs_hook=unique_object,parse_constant=lambda v: (_ for _ in ()).throw(ValueError('비유한 JSON 숫자')))
                if not isinstance(rows,list): raise DerivedValidationError('JSON은 관측값 객체의 배열이어야 합니다.')
            else:
                reader=csv.DictReader(io.StringIO(text,newline=''),delimiter='\t' if suffix=='tsv' else ',')
                if reader.fieldnames!=FIELDS: raise DerivedValidationError('필수 열은 period,value,unit,frequency,source 순서입니다.')
                rows=[]
                for row in reader:
                    if len(rows)>=6000: raise DerivedValidationError('최대 6,000행까지 사용할 수 있습니다.')
                    rows.append(row)
        if not rows or len(rows)>6000: raise DerivedValidationError('파일은 1~6,000개의 관측값을 포함해야 합니다.')
        normalized=[]
        for row in rows:
            if not isinstance(row,dict) or set(row)!=set(FIELDS) or any(isinstance(v,(bool,list,dict)) for v in row.values()):
                raise DerivedValidationError('각 관측값은 필수 필드 5개와 단일 값으로 구성해 주세요.')
            normalized.append({k:'' if row[k] is None else str(row[k]) for k in FIELDS})
        return raw, normalized, suffix
    except DerivedValidationError: raise
    except (ValueError, TypeError, KeyError, UnicodeError, zipfile.BadZipFile, OSError, ImportError, SyntaxError, RuntimeError, zlib.error) as exc:
        raise DerivedValidationError('파일 형식이 올바르지 않습니다. 제공된 양식과 필수 열을 확인해 주세요.') from exc
