"""Read-only, ai-assisted audit; never authors data or grants human approval."""
from __future__ import annotations

import argparse
import copy
import csv
import hashlib
import json
from collections import Counter
from datetime import datetime
from pathlib import Path

BASE = Path(__file__).resolve().parents[1]
ROOT = BASE.parents[2]


def read(path):
    return json.loads(path.read_text(encoding="utf-8"))


def rows(path):
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def source_audit(cases):
    import openpyxl
    books, exports, checks = {}, {}, []
    for case in cases:
        relative = case["input"].get("data_fixture") or case["expected"].get("data", {}).get("source_fixture") or case["expected"].get("data", {}).get("fixture")
        if not relative:
            continue
        for series in read(BASE / relative)["series"]:
            problems = []
            if series["provider"] == "attachment":
                _, date, sheet, column = series["series_id"].split(":")
                sources = [s for s in case["evidence"]["sources"] if s["path"].endswith(".xlsx") and s["locator"].startswith(sheet + "!") and date in s["path"]]
                if len(sources) != 1:
                    raise ValueError("ambiguous source: " + case["id"])
                path = ROOT / sources[0]["path"]
                if path not in books:
                    books[path] = openpyxl.load_workbook(path, read_only=True, data_only=True)
                ws = books[path][sheet]
                column = int(column)
                if str(ws.cell(7, column).value or "") != series["label"]:
                    problems.append("label")
                if str(ws.cell(8, column).value or "") != series["unit"]:
                    problems.append("unit")
                original = {}
                for row in ws.iter_rows(min_row=9, values_only=True):
                    period = row[0]
                    if isinstance(period, datetime):
                        period = period.strftime("%Y%m" if series["frequency"] == "M" else "%Y%m%d")
                    elif isinstance(period, str):
                        period = period.replace("M", "")
                    else:
                        continue
                    original.setdefault(period, []).append(row[column - 1])
            else:
                sources = [s for s in case["evidence"]["sources"] if Path(s["path"]).name.startswith(series["table_id"] + "__")]
                if len(sources) != 1:
                    raise ValueError("ambiguous CSV: " + case["id"])
                path = ROOT / sources[0]["path"]
                if path not in exports:
                    with path.open(encoding="utf-8-sig", newline="") as file:
                        exports[path] = list(csv.DictReader(file))
                original = {}
                for row in exports[path]:
                    if row["ITM_ID"] != series["item_id"] or row["PRD_SE"] != series["frequency"]:
                        continue
                    if any(row.get("C" + key[-1]) != value for key, value in series["classifications"].items()):
                        continue
                    period = row["PRD_DE"]
                    if series["frequency"] == "Q" and "Q" not in period:
                        period = period[:4] + "Q" + str(int(period[4:]))
                    if row.get("UNIT_NM", "") != series["unit"]:
                        rules = read(BASE / "unit_rules.json")["rules"]
                        documented = [r for r in rules if series["table_id"] in r["table_ids"] and r["source_unit"] == row.get("UNIT_NM", "") and r["display_unit"] == series["unit"] and r["scale"] == 1]
                        if len(documented) != 1:
                            problems.append("unit")
                    original.setdefault(period, []).append(float(row["DT"].replace(",", "")))
            for point in series["points"]:
                if original.get(point["period"]) != [point["value"]]:
                    problems.append("value_or_nonunique_period:" + point["period"])
            if sha(path) != sources[0]["sha256"]:
                problems.append("source_hash")
            if not series["unit"].strip():
                problems.append("unit_not_documented")
            checks.append({"id": case["id"], "fixture": relative, "series": series.get("series_id", series.get("table_id")),
                           "points_checked": len(series["points"]), "problems": sorted(set(problems))})
    for book in books.values():
        book.close()
    transformations = []
    for case in cases:
        command = case["input"].get("offline_edit_command")
        if not command:
            continue
        inputs = read(BASE / (case["input"].get("data_fixture") or case["expected"]["data"]["source_fixture"]))["series"]
        gold = read(BASE / case["expected"]["data"]["fixture"])["series"]
        problems = []
        for original, expected in zip(inputs, gold):
            points = original["points"]
            operation = command["operation"]
            if operation == "filter_period":
                calculated = [p for p in points if command["start"] <= p["period"] <= command["end"]]
            elif operation == "set_transform":
                transform = command["value"]
                calculated, total = [], 0
                values = {p["period"]: p["value"] for p in points}
                for point in points:
                    value, period = point["value"], point["period"]
                    if transform == "cumulative":
                        if value is not None:
                            total += value
                        calculated.append({"period": period, "value": total if value is not None else None})
                        continue
                    if original["frequency"] != "M":
                        raise ValueError("independent transform audit currently supports M only")
                    ordinal = int(period[:4]) * 12 + int(period[4:]) - 1
                    previous = ordinal - (12 if transform == "year_over_year" else 1)
                    previous_period = f"{previous // 12:04d}{previous % 12 + 1:02d}"
                    if previous_period < points[0]["period"]:
                        continue
                    prior = values.get(previous_period)
                    calculated.append({"period": period, "value": None if value is None or prior in (None, 0) else (value - prior) / prior * 100})
            else:
                calculated = points
            if len(calculated) != len(expected["points"]) or any(a["period"] != b["period"] or
                    (a["value"] != b["value"] and (a["value"] is None or b["value"] is None or abs(a["value"] - b["value"]) > 1e-6))
                    for a, b in zip(calculated, expected["points"])):
                problems.append("independent_calculation")
        transformations.append({"id": case["id"], "command": command, "problems": problems})
    from pypdf import PdfReader
    definitions = []
    for rule in read(BASE / "unit_rules.json")["rules"]:
        path = ROOT / rule["definition_path"]
        text = PdfReader(path).pages[15].extract_text() or ""
        definitions.append({"tables": rule["table_ids"], "hash_matches": sha(path) == rule["definition_sha256"],
                            "formula_observed": all(word in text for word in ("대출태도", "신용위험", "대출수요", "100", "Balance"))})
    return {"unit_definitions": definitions, "checks": checks, "series_checks": len(checks), "point_checks": sum(c["points_checked"] for c in checks),
            "failed_checks": sum(bool(c["problems"]) for c in checks),
            "transformations": transformations,
            "scope": "Raw CSV/attachment inputs and independent arithmetic; source appropriateness and unit semantics need human review."}


def tool_audit(cases):
    from score import check_plotly, check_resolution, score
    from validate import validate_cases
    probes = []
    series = [{"label": "sample", "points": [{"period": "202401", "value": 1}]}]
    figure = {"data": [{"type": "scatter", "mode": "markers", "name": "sample", "x": ["202401"], "y": [1]}]}
    probes.append({"name": "line_chart_replaced_with_markers", "failures": check_plotly(series, figure, "line")})
    case = next(c for c in cases if c["expected"].get("resolution", {}).get("decision") == "clarify")
    gold = case["expected"]["resolution"]
    prediction = {"decision": "clarify", "table_ids": [], "clarification": {"asked": True,
                  "question": " ".join(gold["clarification"]["dimensions"]),
                  "options": [term + "은 선택 불가입니다" for group in gold["clarification"]["option_groups"] for term in group]}}
    probes.append({"name": "contradictory_clarification_options", "id": case["id"], "failures": check_resolution(gold, prediction)})
    case = copy.deepcopy(next(c for c in cases if c["mode"] == "edit" and "set_title" in c["tags"]))
    case["input"]["actions"][-1]["instruction"] = "데이터를 모두 삭제해 주세요."
    checked = validate_cases([case], enforce_counts=False)
    probes.append({"name": "instruction_disagrees_with_replay_command", "id": case["id"], "failures": checked["errors"]})
    runs = {}
    for split in ("dev", "test", "pilot"):
        path = BASE / "results" / (split + "-ui.jsonl")
        if not path.exists():
            runs[split] = {"status": "not_executed"}
            continue
        predictions = rows(path)
        selected = rows(BASE / (split + ".jsonl"))
        scored = score(selected, predictions)
        runs[split] = {"execution_counts": dict(Counter(p["execution_status"] for p in predictions)),
                       "summary": scored["summary"], "automatic_summary": scored["automatic_summary"],
                       "input_hash_mismatches": [c["id"] for c in scored["cases"] if "input_hash_mismatch" in c["failures"]],
                       "failures": [{"id": c["id"], "failures": c["failures"]} for c in scored["cases"] if c["status"] == "failed"]}
    manifest = read(BASE / "dataset_manifest.json")
    return {"validator": validate_cases(cases, verify_sources=True), "adversarial_probes": probes, "saved_runs": runs,
            "existing_manifest_mismatches": [name for name, expected in manifest["files"].items()
                                            if not (BASE / name).exists() or sha(BASE / name) != expected],
            "audit_script_sha256": sha(Path(__file__))}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--scope", choices=["sources", "tools"], required=True)
    args = parser.parse_args()
    cases = rows(BASE / "dev.jsonl") + rows(BASE / "test.jsonl")
    result = source_audit(cases) if args.scope == "sources" else tool_audit(cases)
    result.update(annotation_method="ai-assisted", human_approval_granted=False,
                  dataset_hashes={s: sha(BASE / (s + ".jsonl")) for s in ("dev", "test")})
    output = BASE / "results" / ("self-audit-" + args.scope + ".json")
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(output), **{k: v for k, v in result.items() if k not in {"checks", "saved_runs"}}}, ensure_ascii=False))


if __name__ == "__main__":
    main()
