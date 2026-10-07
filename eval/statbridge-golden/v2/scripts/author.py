"""Reproduce minimal public fixtures from named originals, never from v1 answers.

Requires openpyxl and pypdf only for authoring, not for offline evaluation.
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import re
from collections import Counter
from datetime import datetime

from common import BASE, ROOT, csv_rows, digest, read_json, read_jsonl, transform_series, write_json, write_jsonl, normalized_unit, unit_sources


class Author:
    def __init__(self):
        self.research = csv_rows(ROOT / "research/그림표-분석/그림표_마스터.csv")
        self.archive = list((ROOT / "research/bok-archive").rglob("*.xlsx"))
        self.catalog_hash = digest(ROOT / "data/processed/bok_table_master.csv")
        self.books, self.cases, self.registry = {}, [], {}
        self.originals = {}

    def workbook(self, date):
        import openpyxl
        paths = [p for p in self.archive if date in str(p) and "그림 원본" in p.name]
        if len(paths) != 1:
            raise ValueError("unique original workbook unavailable: " + date)
        path = paths[0]
        if path not in self.books:
            self.books[path] = openpyxl.load_workbook(path, data_only=True, read_only=True)
        return path, self.books[path]

    def source(self, path, locator, extraction, publisher):
        relative = path.relative_to(ROOT).as_posix()
        self.originals[relative] = digest(path)
        return {"path": relative, "sha256": self.originals[relative], "locator": locator,
                "extraction": extraction, "publisher": publisher}

    def sheet_evidence(self, date, sheet, cell_range="A4:E20"):
        path, book = self.workbook(date)
        matches = [r for r in self.research if r["발간일"] == date and
                   re.search(r"xlsx:" + re.escape(sheet) + r"(?:;|$)", r["근거"])]
        if len(matches) != 1:
            raise ValueError("unique research sheet ID unavailable: " + date + "/" + sheet)
        worksheet = book[sheet]
        title = str(worksheet.cell(4, 1).value)
        if matches[0]["제목"].replace(" ", "") not in title.replace(" ", ""):
            raise ValueError("research title disagrees with original sheet")
        return {"research_result_ids": [matches[0]["결과물ID"]],
                "sources": [self.source(path, sheet + "!" + cell_range,
                    "Read numeric cells with data_only; row 7 labels, row 8 units, column A periods. " +
                    "Retain blank numeric cells as null. No interpolation, inferred purpose, or chart-type labels.",
                    str(worksheet.cell(5, 1).value))],
                "catalog_snapshot": self.catalog_hash, "relationship": "attachment_subset"}

    def subject_evidence(self, split, subject):
        # Explicitly related analytical subjects, not claims of chart-value reproduction.
        if subject in {"money", "rate"}:
            date = "2026-09-10" if split == "dev" else "2026-03-12"
            path, book = self.workbook(date)
            word = "M2 증가율" if subject == "money" else "은행 여수신금리"
            sheets = [s for s in book.sheetnames if word.replace(" ", "") in str(book[s].cell(4, 1).value).replace(" ", "")]
            if len(sheets) != 1:
                raise ValueError("subject figure unavailable: " + word)
            evidence = self.sheet_evidence(date, sheets[0], "A4:K8")
        elif subject == "household":
            date = "2023-12-28" if split == "dev" else "2024-12-24"
            paths = [p for p in self.archive if date in str(p)]
            if len(paths) != 1:
                raise ValueError("financial stability workbook missing")
            import openpyxl
            path = paths[0]
            book = openpyxl.load_workbook(path, data_only=True, read_only=True)
            matches = [r for r in self.research if r["발간일"] == date and r["제목"] == "가계신용" and "xlsx:" in r["근거"]]
            if len(matches) != 1:
                raise ValueError("household research locator missing")
            sheet = re.findall(r"xlsx:([^;]+)", matches[0]["근거"])[0]
            if sheet not in book.sheetnames:
                raise ValueError("household worksheet missing")
            evidence = {"research_result_ids": [matches[0]["결과물ID"]],
                        "sources": [self.source(path, sheet + "!A1:F12",
                            "Verify the named household-credit worksheet; related local KOSIS reconstruction, not copied report values.", "한국은행")],
                        "catalog_snapshot": self.catalog_hash}
        else:
            # PDFs give subject evidence only. Exact values come from separately hashed local exports.
            specs = {
                ("dev", "gdp"): ("BOK-20231130-4a0e2a7937e7-그림-9d1d09d45cf8e0d2", 31, "GDP"),
                ("test", "gdp"): ("BOK-20240222-9ce98d16ec74-그림-788efc7e5ec2069c", 23, "GDP"),
                ("test", "balance"): ("BOK-20241128-e4915bc41ec3-그림-1da53612c381cef6", None, "경상수지"),
                ("test", "lending"): ("BOK-20241128-e4915bc41ec3-그림-3edd55051744d771", None, "대출태도")}
            result_id, page, token = specs[(split, subject)]
            row = next(r for r in self.research if r["결과물ID"] == result_id)
            filename = row["근거"].split(";")[0]
            paths = list((ROOT / "research/bok-archive").rglob(filename))
            if len(paths) != 1:
                raise ValueError("unique PDF missing: " + filename)
            from pypdf import PdfReader
            reader = PdfReader(paths[0])
            if page is None:
                # Use the recorded PDF page only after checking actual original text.
                page = int(re.search(r"PDF p\.(\d+)", row["근거"])[1])
            text = (reader.pages[page - 1].extract_text() or "").replace(" ", "")
            if token not in text:
                raise ValueError("PDF subject not observed at source locator: " + result_id)
            evidence = {"research_result_ids": [result_id], "sources": [self.source(paths[0],
                "PDF page " + str(page) + ", figure " + row["번호"],
                "Verify printed subject on named page; related local KOSIS reconstruction, not report-value or forecast reproduction.",
                row["사용 데이터·출처"])], "catalog_snapshot": self.catalog_hash}
        evidence["relationship"] = "purpose_reconstruction"
        return evidence

    def attachment(self, date, sheet, columns, first=9, count=12, frequency=None):
        path, book = self.workbook(date)
        ws = book[sheet]
        series = []
        for column in columns:
            label, unit = str(ws.cell(7, column).value or ""), str(ws.cell(8, column).value or "")
            points = []
            for row in range(first, first + count):
                period = ws.cell(row, 1).value
                if isinstance(period, datetime):
                    next_period = ws.cell(first+1, 1).value
                    inferred = "W" if isinstance(next_period, datetime) and (next_period-ws.cell(first, 1).value).days == 7 else "D"
                    freq = frequency or ("M" if period.day == 1 and all(isinstance(ws.cell(r, 1).value, datetime) and ws.cell(r, 1).value.day == 1 for r in range(first, first+count)) else inferred)
                    period = period.strftime("%Y%m" if freq == "M" else "%Y%m%d")
                elif isinstance(period, str) and re.fullmatch(r"\d{4}[MQ][0-9]{1,2}", period):
                    freq = "M" if "M" in period else "Q"
                    period = period.replace("M", "") if freq == "M" else period
                else:
                    raise ValueError("unsupported attachment period: " + str(period))
                value = ws.cell(row, column).value
                if value is not None and not isinstance(value, (int, float)):
                    raise ValueError("non-numeric original cell: " + date + "/" + sheet + "/" + str(row) + "/" + str(column) + ": " + repr(value))
                points.append({"period": period, "value": value})
            series.append({"provider": "attachment", "series_id": "attachment:" + date + ":" + sheet + ":" + str(column),
                           "label": label, "unit": unit, "frequency": freq, "points": points})
        from openpyxl.utils import get_column_letter
        return series, self.sheet_evidence(date, sheet, "A4:" + get_column_letter(max(columns)) + str(first+count-1))

    def local(self, table_id, selector, split):
        paths = list((ROOT / "data/tables").glob(table_id + "__*.csv"))
        if len(paths) != 1:
            raise ValueError("unique local export missing: " + table_id)
        rows = csv_rows(paths[0])
        frequency = "Q" if table_id.startswith(("DT_151", "DT_200", "DT_514")) else "M"
        normalize = lambda text: re.sub(r"[\s,]+", "", text)
        names = sorted({r["C1_NM"] for r in rows if normalize(selector) in normalize(r["C1_NM"]) and r["PRD_SE"] == frequency})
        exact = [n for n in names if normalize(n) == normalize(selector)]
        name = exact[0] if exact else names[0] if len(names) == 1 else None
        if name is None:
            raise ValueError("ambiguous local selector: " + table_id + "/" + selector + " " + str(names))
        selected = [r for r in rows if r["C1_NM"] == name and r["PRD_SE"] == frequency]
        start, end = (("202301", "202312") if split == "dev" else ("202401", "202412")) if frequency == "M" else (("2023Q1", "2023Q4") if split == "dev" else ("2024Q1", "2024Q4"))
        expected_count = 12 if frequency == "M" else 4
        if table_id.startswith("DT_101"):
            start, end = ("200301", "200312") if split == "dev" else ("200401", "200409")
            expected_count = 12 if split == "dev" else 9
        def period(r):
            raw = r["PRD_DE"]
            return raw[:4] + "Q" + str(int(raw[-2:])) if frequency == "Q" and "Q" not in raw else raw
        selected = [r for r in selected if start <= period(r) <= end]
        if len(selected) != expected_count:
            raise ValueError("incomplete local period: " + table_id + " " + str([period(r) for r in selected]))
        first = selected[0]
        s = {"provider": "kosis", "table_id": table_id, "item_id": first["ITM_ID"],
             "classifications": {"objL1": first["C1"]}, "label": name, "frequency": frequency,
             "unit": normalized_unit(table_id, first.get("UNIT_NM", "")), "points": sorted([{"period": period(r), "value": float(r["DT"].replace(",", ""))} for r in selected], key=lambda p: p["period"])}
        self.registry[table_id + "|" + first["C1"]] = {k: v for k, v in s.items() if k != "points"}
        source = self.source(paths[0], "CSV: ITM_ID=" + first["ITM_ID"] + ", C1=" + first["C1"] +
                             ", PRD_SE=" + frequency + ", PRD_DE=" + start + ".." + end,
                             "Filter exact item/classification/frequency and inclusive period; numeric DT only; no scaling/interpolation.", "한국은행 KOSIS local export")
        return s, source

    def add(self, split, mode, query, evidence, tags, series=None, input_series=None, decision="resolved",
            sets=None, dimensions=(), options=(), chart="line", command=None, instruction=None, prior=()):
        rid = "SBV2-" + str(len(self.cases) + 1).zfill(4)
        source_group = hashlib.sha256("|".join(sorted(evidence["research_result_ids"])).encode()).hexdigest()[:12]
        case = {"id": rid, "schema_version": "2.0", "split": split, "group_id": split + "-" + source_group,
                "mode": mode, "tags": tags,
                "input": {"query": query, "reference_date": "2026-10-07", "prior_turns": list(prior), "actions": []},
                "expected": {}, "evidence": copy.deepcopy(evidence),
                "review": {"status": "ready_for_review", "reviewer": None, "human_approved": False,
                           "annotation_method": "ai-assisted-source-grounded"}}
        if mode in {"discovery", "e2e"}:
            case["expected"]["resolution"] = {"decision": decision, "acceptable_table_sets": sets or [],
                "clarification": {"required": decision == "clarify", "dimensions": list(dimensions), "option_groups": list(options)}}
            if decision == "resolved":
                case["expected"]["resolution"]["series_selection"] = [
                    {k: v for k, v in s.items() if k != "points"}
                    for s in self.registry.values() if s["table_id"] in (sets or [[]])[0]]
        if series is not None:
            relative = "fixtures/" + rid + ".json"
            write_json(BASE / relative, {"series": series})
            case["expected"]["data"] = {"fixture": relative, "fixture_sha256": digest(BASE / relative),
                "absolute_tolerance": 0.000001 if command else 0.0000001, "period_policy": "exact", "missing_policy": "preserve"}
            claims = "claims/" + rid + ".json"
            # Keep authoring runtime independent of the validator's jsonschema dependency.
            facts = []
            from common import identity
            for s in series:
                valid = [p for p in s["points"] if p["value"] is not None]
                fact = {"identity": list(identity(s)), "unit": s["unit"], "frequency": s["frequency"],
                        "missing_periods": [p["period"] for p in s["points"] if p["value"] is None]}
                if valid:
                    fact.update(first=valid[0], last=valid[-1], minimum=min(valid, key=lambda p: p["value"]),
                        maximum=max(valid, key=lambda p: p["value"]), absolute_change=valid[-1]["value"]-valid[0]["value"], change_formula="last - first")
                facts.append(fact)
            write_json(BASE / claims, {"fixture_sha256": digest(BASE / relative), "facts": facts})
            layout_series = input_series if input_series is not None else series
            layout = "separate" if len({s["unit"] for s in layout_series}) > 1 or len({s["frequency"] for s in layout_series}) > 1 else "combined"
            case["expected"]["output"] = {"acceptable_chart_types": [chart], "layout": layout, "claims": claims}
            case["input"]["output_request"] = {"chart_type": chart, "layout": "combined"}
            if mode == "e2e":
                points = series[0]["points"]
                start, end = points[0]["period"], points[-1]["period"]
                first_month = (int(start[-1])*3-2) if "Q" in start else int(start[4:6])
                last_month = int(end[-1])*3 if "Q" in end else int(end[4:6])
                import calendar
                start_date = start[:4] + "-" + str(first_month).zfill(2) + "-01"
                end_date = end[:4] + "-" + str(last_month).zfill(2) + "-" + str(calendar.monthrange(int(end[:4]), last_month)[1])
                case["input"]["actions"] = [{"kind": "confirm_period", "start": start_date, "end": end_date},
                                            {"kind": "configure_output", "chart_type": chart, "layout": "combined"}]
            if mode in {"output", "edit"}:
                attached = "fixtures/" + rid + "-input.json"
                write_json(BASE / attached, {"series": input_series if input_series is not None else series})
                case["input"]["data_fixture"] = attached
                case["input"]["data_fixture_sha256"] = digest(BASE / attached)
            if command:
                case["input"]["offline_edit_command"] = command
                case["input"]["actions"].append({"kind": "edit", "instruction": instruction})
                state = {}
                if command["operation"] == "set_title": state["title"] = command["value"]
                if command["operation"] == "set_legend": state["show_legend"] = command["value"]
                if command["operation"] == "set_transform": state["transform"] = command["value"]
                if command["operation"] == "filter_period": state.update(period_start=command["start"], period_end=command["end"])
                case["expected"]["output"].update(state=state, preserve_data=True)
        self.cases.append(case)
        return case


def build():
    a = Author()
    specs = {
        "dev": [
            ("DT_101Y001", "M2(말잔, 계절조정계열)", "money", "통화 및 유동성지표(구)의 M2 상품별 구성내역 중 계절조정 말잔 M2"),
            ("DT_101Y002", "M2(말잔, 원계열)", "money", "통화 및 유동성지표(구)의 M2 상품별 구성내역 중 원계열 말잔 M2"),
            ("DT_121Y002", "저축성수신(금융채 제외)", "rate", "예금은행 신규취급액 기준 저축성수신 금리(금융채 제외)"),
            ("DT_121Y006", "대출평균", "rate", "예금은행 신규취급액 기준 대출평균 금리"),
            ("DT_151Y001", "가계신용", "household", "기관별 가계신용의 가계신용 전체 잔액"),
            ("DT_200Y107", "국내총생산에 대한 지출", "gdp", "지출에 의한 국내총생산 계절조정 명목"),
            ("DT_200Y108", "민간", "gdp", "지출에 의한 국내총생산의 민간소비 계절조정 실질")],
        "test": [
            ("DT_101Y003", "M2(평잔, 계절조정계열)", "money", "통화 및 유동성지표(구)의 M2 상품별 구성내역 중 계절조정 평잔 M2"),
            ("DT_101Y004", "M2(평잔, 원계열)", "money", "통화 및 유동성지표(구)의 M2 상품별 구성내역 중 원계열 평잔 M2"),
            ("DT_121Y013", "저축예금", "rate", "예금은행 잔액 기준 저축예금 금리"),
            ("DT_121Y015", "가계대출", "rate", "예금은행 잔액 기준 가계대출 금리"),
            ("DT_151Y004", "주택관련대출", "household", "용도별 가계신용 중 주택관련대출 잔액"),
            ("DT_200Y109", "건설투자", "gdp", "지출에 의한 국내총생산의 건설투자 원계열 명목"),
            ("DT_200Y110", "재화와 서비스의 수출", "gdp", "지출에 의한 국내총생산의 재화와 서비스 수출 원계열 실질"),
            ("DT_301Y017", "경상수지", "balance", "계절조정 경상수지"),
            ("DT_514Y001", "국내은행-차주가중종합지수", "lending", "대출태도지수 국내은행 차주가중종합지수"),
            ("DT_514Y002", "국내은행-차주가중종합지수", "lending", "신용위험지수 국내은행 차주가중종합지수"),
            ("DT_514Y003", "국내은행-차주가중종합지수", "lending", "대출수요지수 국내은행 차주가중종합지수")]}
    cached = {}
    for split, rows in specs.items():
        year = "2023" if split == "dev" else "2024"
        for table, selector, subject, metric in rows:
            series, source = a.local(table, selector, split)
            evidence = a.subject_evidence(split, subject)
            evidence["sources"].append(source)
            evidence["sources"].extend(unit_sources(table))
            cached[(split, table)] = (series, evidence, metric)
            unit = "분기" if series["frequency"] == "Q" else "월"
            interval = ("2003년 1월~12월" if split == "dev" else "2004년 1월~9월") if table.startswith("DT_101") else year+"년 1월부터 12월까지"
            a.add(split, "e2e", interval+" "+metric+"를 "+unit+"별 선그래프로 보여주세요.",
                  evidence, ["single", subject], [series], sets=[[table]])
    for split in ("dev", "test"):
        year = "2023" if split == "dev" else "2024"
        pairs = ([("DT_101Y001", "DT_101Y002"), ("DT_121Y002", "DT_121Y006"), ("DT_200Y107", "DT_200Y108")]
                 if split == "dev" else [("DT_101Y003", "DT_101Y004"), ("DT_121Y013", "DT_121Y015"), ("DT_514Y002", "DT_514Y003")])
        for left, right in pairs:
            l, ev, lm = cached[(split, left)]
            r, rev, rm = cached[(split, right)]
            evidence = copy.deepcopy(ev)
            evidence["research_result_ids"] = sorted(set(ev["research_result_ids"] + rev["research_result_ids"]))
            evidence["sources"] = [*ev["sources"], *[s for s in rev["sources"] if s not in ev["sources"]]]
            interval = ("2003년 1월~12월" if split == "dev" else "2004년 1월~9월") if left.startswith("DT_101") else year+"년"
            a.add(split, "e2e", interval+" "+lm+"와 "+rm+"를 각각 원래 단위로 비교하는 선그래프를 만들어주세요.",
                  evidence, ["multi", "comparison"], [l, r], sets=[[left, right]])
        follow_tables = (["DT_121Y006", "DT_151Y001"] if split == "dev" else ["DT_301Y017", "DT_121Y015", "DT_200Y110", "DT_514Y003"])
        for table in follow_tables:
            s, ev, metric = cached[(split, table)]
            prior = [{"role": "user", "text": str(int(year)-1)+"년 "+metric+"의 같은 분류와 주기 자료를 조회해 주세요."},
                     {"role": "assistant", "text": str(int(year)-1)+"년 "+metric+" 조회에서 기간은 1월~12월, 원래 단위, 계열 하나로 확정했습니다."}]
            a.add(split, "e2e", "앞서 조회한 "+metric+"의 분류·주기는 그대로 두고 기간만 "+year+"년 1월~12월로 바꾸어 막대그래프로 보여주세요.",
                  ev, ["followup", "period_change"], [s], sets=[[table]], chart="bar", prior=prior)
        resolved_tables = [r[0] for r in specs[split]][:5 if split == "dev" else 7]
        for table in resolved_tables:
            _, ev, metric = cached[(split, table)]
            interval = ("2003년 1월~12월" if split == "dev" else "2004년 1월~9월") if table.startswith("DT_101") else year+"년"
            a.add(split, "discovery", interval+" "+metric+"를 조회할 통계표를 찾아주세요. 수치는 아직 조회하지 마세요.",
                  ev, ["selectable"], sets=[[table]])
    clarify_specs = {
        "dev": [("money", "2023년 M2를 보고 싶습니다. 말잔인지 평잔인지는 정하지 않았습니다.", ["잔액 기준"], [["말잔", "평잔"]]),
                ("rate", "2023년 예금은행 금리를 찾아주세요. 예금금리인지 대출금리인지는 아직 정하지 않았어요.", ["금리"], [["수신", "대출"]]),
                ("gdp", "2023년 분기별 GDP를 보고 싶습니다. 명목과 실질 중 어떤 기준인지 먼저 물어봐 주세요.", ["기준"], [["명목", "실질"]])],
        "test": [("money", "2024년 M2 평잔을 보고 싶습니다. 계절조정 여부는 정하지 않았습니다.", ["계절"], [["계절조정", "원계열"]]),
                 ("rate", "2024년 예금은행 대출금리를 찾습니다. 신규취급액과 잔액 중 기준을 고르지 않았습니다.", ["기준"], [["신규취급", "잔액"]]),
                 ("household", "2024년 가계대출을 보고 싶습니다. 기관별인지 용도별인지부터 선택하고 싶어요.", ["분류"], [["기관", "용도"]]),
                 ("gdp", "2024년 실질 GDP를 찾습니다. 원계열인지 계절조정인지 먼저 확인해 주세요.", ["계절"], [["원계열", "계절조정"]]),
                 ("lending", "2024년 국내은행 대출행태를 보려는데 태도·신용위험·수요 중 지표를 정하지 않았습니다.", ["지표"], [["태도", "신용위험", "수요"]])]}
    for split, rows in clarify_specs.items():
        for subject, query, dims, opts in rows:
            a.add(split, "discovery", query, a.subject_evidence(split, subject), ["clarification"], decision="clarify", dimensions=dims, options=opts)
    # External subjects are present in research, but no attachment is provided here.
    for split, date, queries in [
        ("dev", "2026-09-10", [("개요1", "2024년 미국과 일본의 10년 국채금리를 한국은행 KOSIS 조회만으로 보여주세요."),
                                ("개요5", "2022년 서울 주택매매가격 상승률을 한국은행 KOSIS 조회만으로 보여주세요.")]),
        ("test", "2026-03-12", [("개요1", "2024년 달러화 DXY 지수를 한국은행 KOSIS 조회만으로 찾아주세요."),
                                 ("개요5", "2022년 비수도권 주택매매가격 상승률을 한국은행 KOSIS 조회만으로 찾아주세요.")])]:
        for sheet, query in queries:
            evidence = a.sheet_evidence(date, sheet)
            evidence["relationship"] = "unsupported_subject"
            a.add(split, "discovery", query, evidence, ["unsupported", "external_without_attachment"], decision="unsupported")
    output_specs = [
        ("dev", "2026-09-10", "개요1", [4, 5], 9, 12, ["external", "comparison"]),
        ("dev", "2026-09-10", "개요3", [2, 4], 9, 12, ["mixed_source", "comparison"]),
        ("dev", "2026-09-10", "개요4", [2, 4, 5], 9, 12, ["mixed_source", "mixed_units"]),
        ("dev", "2026-09-10", "개요5", [2, 3, 4], 9, 12, ["external", "growth_rate"]),
        ("dev", "2026-09-10", "개요8", [2, 3], 9, 24, ["mixed_units", "transformation"]),
        ("test", "2026-03-12", "개요1", [2, 3], 9, 12, ["external", "comparison"]),
        ("test", "2026-03-12", "개요3", [2, 3], 9, 16, ["mixed_source", "missing"]),
        ("test", "2026-03-12", "개요4", [2, 3], 33, 12, ["mixed_source", "comparison"]),
        ("test", "2026-03-12", "개요5", [2, 4], 9, 12, ["external", "negative_values"]),
        ("test", "2026-03-12", "개요6", [2, 3, 4, 5], 9, 12, ["mixed_source", "transformation"]),
        ("test", "2026-03-12", "개요7", [2], 9, 12, ["daily", "step_series"]),
        ("test", "2026-03-12", "개요8", [2], 33, 24, ["mixed_frequency"])]
    for split, date, sheet, cols, first, count, tags in output_specs:
        s, ev = a.attachment(date, sheet, cols, first, count)
        command, instruction = None, None
        if "mixed_frequency" in tags:
            other, other_ev = a.attachment(date, "개요7", [2], 9, 12)
            s += other
            ev["research_result_ids"] += other_ev["research_result_ids"]
            ev["sources"] += other_ev["sources"]
        if "transformation" in tags:
            command = {"operation": "set_transform", "kind": "DATA_EDIT", "value": "cumulative" if sheet == "개요6" else "growth_rate"}
            instruction = "첨부된 값으로 누적합을 계산해 주세요." if sheet == "개요6" else "첨부 계열을 전기 대비 증감률(%)로 바꿔 주세요."
        output = transform_series(s, command) if command else s
        period_description = ", ".join(str(x["points"][0]["period"])+"~"+str(x["points"][-1]["period"]) for x in s)
        a.add(split, "output", "첨부한 "+sheet+" 자료("+period_description+")의 모든 계열을 원래 단위·주기로 선그래프에 표시하고 첫 값, 마지막 값, 변화와 결측을 설명해주세요.",
              ev, tags, output, input_series=s, command=command, instruction=instruction)
    edit_specs = [
        ("dev", "2025-03-13", "set_title", "금융중개지원대출: 한도와 금리", "제목을 '금융중개지원대출: 한도와 금리'로 변경해 주세요."),
        ("dev", "2025-03-13", "set_legend", False, "범례를 숨겨 주세요. 데이터와 출처는 유지해 주세요."),
        ("dev", "2025-03-13", "set_transform", "growth_rate", "전기 대비 증감률(%)로 변경해 주세요."),
        ("test", "2025-09-11", "set_title", "지원대출 한도·금리 추이", "제목을 '지원대출 한도·금리 추이'로 변경해 주세요."),
        ("test", "2025-09-11", "filter_period", None, "2015년 1월부터 12월까지로 표시 기간을 좁혀 주세요."),
        ("test", "2025-09-11", "set_transform", "year_over_year", "전년 같은 달 대비 증감률(%)로 변경해 주세요.")]
    for split, date, op, value, instruction in edit_specs:
        first = 9 if split == "dev" else 57
        s, ev = a.attachment(date, "개요8", [2, 3] if split == "dev" else [2], first, 24, "M")
        command = {"operation": op, "kind": "DATA_EDIT" if op in {"set_transform", "filter_period"} else "STYLE_EDIT"}
        if op == "filter_period": command.update(start="201501", end="201512")
        else: command["value"] = value
        transformed = transform_series(s, command)
        interval = "2010년 1월~2011년 12월" if split == "dev" else "2014년 1월~2015년 12월"
        a.add(split, "edit", "첨부한 금융중개지원대출 자료의 "+interval+" 선그래프를 먼저 생성해 주세요.",
              ev, [op, "state_preservation"], transformed, input_series=s, command=command, instruction=instruction)
    # Connected components: every derivative of a source or series-period belongs
    # to one group, including a mixed attachment that joins two original figures.
    parents = {c["id"]: c["id"] for c in a.cases}
    def find(x):
        while parents[x] != x:
            x = parents[x]
        return x
    seen = {}
    from common import identity
    for c in a.cases:
        keys = [("research", r) for r in c["evidence"]["research_result_ids"]]
        relative = c["input"].get("data_fixture") or c["expected"].get("data", {}).get("fixture")
        if relative:
            for s in read_json(BASE / relative)["series"]:
                core = identity(s) if s["provider"] == "kosis" else (
                    "attachment_subject", re.sub(r"\s+", "", s["label"]), re.sub(r"\s+", "", s["unit"]), s["frequency"])
                keys += [(core, p["period"]) for p in s["points"]]
        for key in keys:
            if key in seen:
                other = seen[key]
                if other["split"] != c["split"]:
                    raise ValueError("source/series-period split leakage: " + str(key))
                parents[find(c["id"])] = find(other["id"])
            seen[key] = c
    for c in a.cases:
        c["group_id"] = c["split"] + "-group-" + find(c["id"])
    write_json(BASE / "series_registry.json", list(a.registry.values()))
    attachments = {}
    for c in a.cases:
        if c["input"].get("data_fixture"):
            for s in read_json(BASE / c["input"]["data_fixture"])["series"]:
                attachments[s["series_id"]] = {k: v for k, v in s.items() if k != "points"}
    write_json(BASE / "attachment_registry.json", list(attachments.values()))
    write_json(BASE / "fixtures/provider.json", {"series": [value[0] for value in cached.values()]})
    candidate_counts = Counter(r["보고서"].split(" ")[0] for r in a.research)
    selected_ids = {rid for c in a.cases for rid in c["evidence"]["research_result_ids"]}
    selected = Counter(r["보고서"].split(" ")[0] for r in a.research if r["결과물ID"] in selected_ids)
    write_json(BASE / "source_inventory.json", {"candidate_report_types": dict(candidate_counts),
        "selected_unique_result_ids_by_report_type": dict(selected), "selected_unique_results": len(selected_ids),
        "selection_note": "Five report types inspected; only originals with explicit locators selected. No quota filling from inferred fields.",
        "excluded_candidates": [{"subject": "Financial support loan rate 2012-2015", "reason": "Original rate cells contain ranges, e.g. 0.5~1.25; excluded from scalar numeric tasks. Loan limits retained."}],
        "originals": a.originals})
    return a.cases


def main():
    if (BASE / "dataset_policy.json").exists():
        raise ValueError("legacy attachment author is archived; public UI revision is maintained by ui_revision.py")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pilot", action="store_true")
    args = parser.parse_args()
    existing = [c for split in ("dev", "test") if (BASE / (split+".jsonl")).exists()
                for c in read_jsonl(BASE / (split+".jsonl"))]
    if any(c["review"]["human_approved"] for c in existing):
        raise ValueError("refusing to overwrite human-approved cases; author a separately reviewed version")
    cases = build()
    # Balanced representative subset, retained in final public data.
    pilot = []
    for mode in ("discovery", "e2e", "output", "edit"):
        rows = [c for c in cases if c["mode"] == mode]
        pilot.extend([rows[0], rows[len(rows)//2], rows[-1]])
    write_jsonl(BASE / "pilot.jsonl", pilot)
    if not args.pilot:
        for split in ("dev", "test"):
            write_jsonl(BASE / (split + ".jsonl"), [c for c in cases if c["split"] == split])
        write_jsonl(BASE / "review/gold-review-template.jsonl", [
            {"id": c["id"], "status": "pending", "reviewer": None, "human_approved": False,
             "question_clear": None, "original_locator_verified": None, "numbers_units_frequency_verified": None,
             "acceptable_answers_verified": None, "split_independence_verified": None,
             "source_relationship": c["evidence"]["relationship"], "notes": ""} for c in cases])
    print({"authored": len(cases), "pilot": len(pilot), "published": 12 if args.pilot else len(cases), "human_approved": 0})


if __name__ == "__main__":
    main()
