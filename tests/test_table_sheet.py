"""Free sheet editing: immutable observations, reversible spans and exports."""
import copy
import io
import json
import pytest
from report_table import TableOptions,build_table,render_table_file
from test_table_layout import CASES

def case():return copy.deepcopy(CASES[0])
def doc(c):return build_table(c['series'],TableOptions.model_validate(c['options']),title=c['title'])
def merge(c,rows=2,cols=2):
    d=doc(c);s=d['layoutSections'][0];body=[r['id'] for r in s['sheetRows'] if r['kind']=='body'];columns=[x['id'] for x in s['columns']][1:1+cols]
    c['options']['layout']['cell_merges']=[{'id':'free','rows':body[:rows],'columns':columns}];return body[:rows],columns

@pytest.mark.parametrize('rows,cols',[(1,2),(2,1),(2,2)])
def test_arbitrary_horizontal_vertical_rectangular_merge_and_split(rows,cols):
    c=case();before=copy.deepcopy(c['series']);original=doc(c);merge(c,rows,cols);d=doc(c)
    cell=next(x for r in d['layoutSections'][0]['sheetRows'] for x in r['cells'] if x['mergeId']=='merge:free')
    assert (cell['rowSpan'],cell['colSpan'])==(rows,cols)
    assert len(cell['text'].splitlines())==rows*cols and c['series']==before
    assert d['layoutSections'][0]['layoutRows']==original['layoutSections'][0]['layoutRows']
    c['options']['layout']['cell_merges']=[];assert doc(c)['layoutSections'][0]['sheetRows']==original['layoutSections'][0]['sheetRows']

def test_split_existing_multirow_header_preserves_every_underlying_cell():
    c=case();d=doc(c);header=next(x for x in d['layoutSections'][0]['sheetRows'][0]['cells'] if x['colSpan']==3)
    c['options']['layout']['split_cells']=[header['mergeId']];updated=doc(c)['layoutSections'][0]['sheetRows'][0]['cells']
    assert len(updated)==len(d['layoutSections'][0]['sheetRows'][0]['cells'])+2
    assert sum(x['text']==header['text'] for x in updated)==1

@pytest.mark.parametrize('kind',['csv','xlsx','pdf','png'])
def test_merged_cells_and_range_colors_export_without_losing_evidence(kind):
    c=case();rr,cc=merge(c);c['options']['layout']['range_formats']=[{'id':'color','rows':rr,'columns':cc,'format':{'background':'#fff2cc','bold':True,'bottom':'medium'}}]
    d=doc(c);blob=render_table_file(d,kind);assert len(blob)>300
    merged=next(x for r in d['layoutSections'][0]['sheetRows'] for x in r['cells'] if x['mergeId']=='merge:free');assert merged['format']['background']=='#fff2cc'
    if kind=='xlsx':
        from openpyxl import load_workbook
        wb=load_workbook(io.BytesIO(blob));ws=wb.worksheets[0]
        assert 'B5:C6' in {str(m) for m in ws.merged_cells.ranges};assert ws['B5'].fill.fgColor.rgb=='00fff2cc'
        assert wb['계산근거'].max_row==37
        assert all(c.data_type!='f' for sheet in wb for row in sheet for c in row)

@pytest.mark.parametrize('patch',[
 {'cell_merges':[{'id':'m','rows':['missing'],'columns':['label']}]},
 {'cell_merges':[{'id':'m','rows':['matrix-row:r0','matrix-row:r2'],'columns':['matrix:h0a']}]},
 {'cell_merges':[{'id':'m','rows':['matrix-row:r0'],'columns':['matrix:h0a','matrix:h0d']}]},
 {'cell_merges':[{'id':'m','rows':['matrix-row:r0'],'columns':['matrix:h0a']}]},
 {'cell_merges':[{'id':'a','rows':['matrix-row:r0'],'columns':['matrix:h0a','matrix:h0b']},{'id':'b','rows':['matrix-row:r0'],'columns':['matrix:h0b','matrix:h0d']}]},
 {'cell_merges':[{'id':'h','rows':['sheet-header:0:0'],'columns':['matrix:h0a','matrix:h0b']}]},
 {'split_cells':['no-such-merge']},
 {'range_formats':[{'id':'f','rows':['missing'],'columns':['label'],'format':{'background':'#fff2cc'}}]},
 {'cell_text':{'matrix-row:r0|matrix:h0a':'99999'}},
 {'cell_text':{'missing|label':'bad'}},
 {'heights':{'sheet-header:0:4':30}},
 {'cell_merges':[{'id':'m','rows':['matrix-row:r0'],'columns':['matrix:h0a','matrix:h0b'],'code':'exec(x)'}]},
])
def test_invalid_or_destructive_edit_is_rejected(patch):
    c=case();before=copy.deepcopy(c['series']);c['options']['layout'].update(patch)
    with pytest.raises(ValueError):doc(c)
    assert c['series']==before

def test_label_edit_and_range_format_restore_with_original_values_unchanged():
    c=case();c['options']['layout']['cell_text']={'matrix-row:r0|label':'원문 대신 표시할 제목'}
    c['options']['layout']['range_formats']=[{'id':'a','rows':['matrix-row:r0'],'columns':['label','matrix:h0a'],'format':{'background':'#e2f0d9','bold':True}}]
    d=doc(c);restored=build_table(c['series'],TableOptions.model_validate(json.loads(json.dumps(d['options']))),title=c['title']);assert restored==d
    row=d['layoutSections'][0]['sheetRows'][2];assert row['cells'][0]['text']=='원문 대신 표시할 제목';assert row['cells'][1]['format']['bold']
    assert row['cells'][1]['text']=='0.8'

def test_merged_cell_projection_survives_pdf_page_boundaries():
    c=copy.deepcopy(CASES[2]);s=doc(c)['layoutSections'][0];body=[r['id'] for r in s['sheetRows'] if r['kind']=='body'];c['options']['layout']['rows_per_page']=5
    # Existing section heading rows are included in the rectangle explicitly.
    selected=[r['id'] for r in s['sheetRows']][3:8]
    c['options']['layout']['cell_merges']=[{'id':'long','rows':selected,'columns':['period:2024S1']}]
    assert render_table_file(doc(c),'pdf').startswith(b'%PDF')


def test_single_header_cell_is_not_offered_as_a_merge():
    d=doc(case())
    for row in d["layoutSections"][0]["sheetRows"]:
        for cell in row["cells"]:
            if cell["rowSpan"]==cell["colSpan"]==1:
                assert not cell["mergeId"]


@pytest.mark.parametrize('style',['none','thin','medium','thick','dotted','dotted_medium','dotted_thick','dashed','dashed_medium','dashed_thick','double'])
def test_individual_border_styles_colors_export_and_restore(style):
    from table_borders import excel_border
    from openpyxl import load_workbook
    c=case();d=doc(c);s=d['layoutSections'][0];row=next(r for r in s['sheetRows'] if r['kind']=='body');col=s['columns'][1]['id']
    c['options']['layout']['range_formats']=[{'id':'edge','rows':[row['id']],'columns':[col],'format':{'top':style,'top_color':'#ff0000','bottom':'medium','bottom_color':'#00aa00','left':'dashed','right':'double'}}]
    d=doc(c);cell=next(x for r in d['layoutSections'][0]['sheetRows'] if r['id']==row['id'] for x in r['cells'] if x['column']==1)
    assert cell['format']['top']==style and cell['format']['top_color']=='#ff0000'
    wb=load_workbook(io.BytesIO(render_table_file(d,'xlsx')));x=wb.worksheets[0].cell(cell['row']+3,cell['column']+1)
    assert x.border.top.style==excel_border(style) and x.border.top.color.rgb=='00ff0000'
    assert x.border.bottom.style=='medium' and x.border.bottom.color.rgb=='0000aa00'
    assert x.border.left.style=='dashed' and x.border.right.style=='double'
    assert build_table(c['series'],TableOptions.model_validate(json.loads(json.dumps(d['options']))),title=c['title'])==d


def test_range_number_format_changes_display_only_and_merged_boundary_edges():
    c=case();rr,cc=merge(c);c['options']['layout']['range_formats']=[{'id':'number','rows':rr,'columns':cc,'format':{'decimals':3,'signed':True}}, {'id':'bottom','rows':rr[-1:],'columns':cc,'format':{'bottom':'thick','bottom_color':'#ff0000'}}]
    d=doc(c);cell=next(x for r in d['layoutSections'][0]['sheetRows'] for x in r['cells'] if x['mergeId']=='merge:free')
    assert cell['format']['bottom']=='thick' and cell['format']['bottom_color']=='#ff0000'
    assert all(len(line.split('.')[-1])==3 for line in cell['text'].splitlines())
    original=copy.deepcopy(c);original['options']['layout']['range_formats']=[];assert d['layoutSections'][0]['layoutRows'][0]['cells']==doc(original)['layoutSections'][0]['layoutRows'][0]['cells']


def test_every_template_has_visible_sheet_style_and_manual_overrides_win():
    from test_bok_table_styles import observations
    fingerprints=[]
    for template in ['default',*[f'BOK-T0{i}' for i in range(1,7)]]:
        d=build_table(observations(),TableOptions(template=template,decimals=1,emphasized_labels=['조회된 소계'],layout={'enabled':True}))
        s=d['layoutSections'][0];fingerprints.append(json.dumps([[c['format'] for c in r['cells']] for r in s['sheetRows']],sort_keys=True))
        assert len(s['headerRows'])==(2 if template in {'BOK-T02','BOK-T03'} else 1)
        assert d['sections'][0]['rows'][0][1:]==[1.234,0,-2.5]
    assert len(set(fingerprints))==7
    d=build_table(observations(),TableOptions(template='BOK-T06',layout={'enabled':True,'base':{'background':'#123456'}}))
    assert all(c['format']['background']=='#123456' for r in d['layoutSections'][0]['sheetRows'] for c in r['cells'])


def test_extended_border_image_and_pdf_export():
    c=case();rr,cc=merge(c);c['options']['layout']['range_formats']=[{'id':'edge','rows':rr,'columns':cc,'format':{'top':'dotted_thick','bottom':'double','left':'dashed_medium','right':'thick','top_color':'#ff0000','vertical_align':'bottom'}}]
    d=doc(c);assert render_table_file(d,'png').startswith(b'\x89PNG');assert render_table_file(d,'pdf').startswith(b'%PDF')


def test_latest_highlight_targets_actual_observation_not_appended_comparison():
    from test_bok_table_styles import observations
    d=build_table(observations(),TableOptions(template='BOK-T02',comparison='period_change',layout={'enabled':True}))
    s=d['layoutSections'][0];row=next(r for r in s['layoutRows'] if not r.get('group'))
    assert row['formats']['period:202502']['background']=='#edf7f3'
    assert row['formats']['legacy_comparison']['background']=='#ffffff'
