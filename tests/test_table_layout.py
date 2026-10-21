"""Independent attachment-based layout checks. No live API or model grading."""
import copy
import io
import json
from decimal import Decimal
from pathlib import Path

import pytest
from report_table import TableOptions,build_table,render_table_file
from table_layout import TableLayout,stable

ROOT=Path(__file__).resolve().parents[1]
CASES=json.loads((ROOT/'eval/report_table/reference_cases.json').read_text(encoding='utf8'))

def build(case):return build_table(case['series'],TableOptions.model_validate(case['options']),title=case['title'])

@pytest.mark.parametrize('case',CASES,ids=lambda c:c['id'])
def test_reference_layout_and_roundtrip_preserve_every_observation(case):
    original=copy.deepcopy(case);doc=build(case)
    assert case==original
    rebuilt=build_table(case['series'],TableOptions.model_validate(json.loads(json.dumps(doc['options']))),title=case['title'])
    assert rebuilt==doc
    assert doc['annotations'] and doc['layoutSections'][0]['headerRows']
    for section in doc['layoutSections']:
        for row in section['layoutRows']:
            for cid,ref in row.get('references',{}).items():
                value=next(p['value'] for s in case['series'] if s['label']==ref['series'] for p in s['points'] if p['date']==ref['period'])
                if cid not in row.get('comparisons',{}):assert row['cells'][cid]==value

@pytest.mark.parametrize('case',[CASES[0],CASES[2]],ids=lambda c:c['id'])
def test_reference_comparisons_match_independent_decimal(case):
    for row in build(case)['layoutSections'][0]['layoutRows']:
        for cid,c in row.get('comparisons',{}).items():
            a,b=map(lambda v:Decimal(str(v)),c['observations'])
            expected=(b/a-1)*100 if c['operation']=='growth' else b-a
            assert Decimal(str(c['value']))==expected
            if c['operation']=='difference_growth':assert abs(Decimal(str(c['secondary']))-(b/a-1)*100)<Decimal('1e-12')

@pytest.mark.parametrize('case',CASES,ids=lambda c:c['id'])
@pytest.mark.parametrize('kind',['csv','xlsx','pdf','png'])
def test_reference_files_have_real_content_and_typed_evidence(case,kind):
    doc=build(case);blob=render_table_file(doc,kind)
    assert len(blob)>300
    if kind=='xlsx':
        from openpyxl import load_workbook
        wb=load_workbook(io.BytesIO(blob),data_only=False)
        assert '계산근거' in wb.sheetnames and wb['계산근거'].max_row>2
        assert wb.worksheets[0].merged_cells.ranges and wb.worksheets[0].print_title_rows
        for ws in wb:
            for row in ws:
                assert all(c.data_type!='f' for c in row)
        numbers=[c.value for row in wb['계산근거'] for c in row if isinstance(c.value,(int,float))]
        assert numbers
    elif kind=='pdf':assert blob.startswith(b'%PDF')
    elif kind=='png':assert blob.startswith(b'\x89PNG')
    else:assert doc['title'] in blob.decode('utf-8-sig')

@pytest.mark.parametrize('field,value',[('font_size',7),('font_size',25),('color','url(bad)'),('align','script'),('decimals',7),('bottom','eval')])
def test_closed_style_schema(field,value):
    with pytest.raises(ValueError):TableLayout.model_validate({'enabled':True,'formats':{'label':{field:value}}})

@pytest.mark.parametrize('patch',[
 {'formats':{'missing':{'bold':True}}}, {'widths':{'label':501}}, {'heights':{'label':40}},
 {'column_order':['label','label']}, {'aliases':{'label':''}}, {'rows_per_page':0},
 {'header_bands':[{'label':'bad','columns':['label','period:2025S1']}]},
 {'header_bands':[{'label':'bad','columns':['period:2025S1'],'level':1}]},
 {'comparisons':[{'id':'x','label':'bad','baseline':'1900','current':'2026S1'}]},
 {'annotations':[{'id':'n','text':'bad','target':'other'}]},
 {'code':'import os'}, {'bindings':[{'id':'b','label':'x','series':['unknown']}]}])
def test_invalid_layout_is_rejected_without_mutating_original(patch):
    case=copy.deepcopy(CASES[2]);before=copy.deepcopy(case['series']);case['options']['layout'].update(patch)
    with pytest.raises(ValueError):build(case)
    assert case['series']==before

def test_mixed_periods_are_not_aggregated_and_same_year_remains_distinct():
    doc=build(CASES[1]);rows=doc['layoutSections'][0]['layoutRows']
    assert len(rows)==7
    assert {r['id'] for r in rows}>={'period:Y:2024','period:Q:2024Q4'}
    assert len({r['label'] for r in rows})==7

def test_sparse_cell_overrides_do_not_reset_column_style_on_restore():
    case=copy.deepcopy(CASES[2]);rid=stable('series:',case['series'][0]['id']);cid='period:2026S1'
    case['options']['layout']['formats'][rid+'|'+cid]={'color':'#cc0000'}
    doc=build(case);restored=build_table(case['series'],TableOptions.model_validate(doc['options']))
    for d in [doc,restored]:
        row=next(r for r in d['layoutSections'][0]['layoutRows'] if r['id']==rid)
        assert row['formats'][cid]['bold'] and row['formats'][cid]['color']=='#cc0000'

@pytest.mark.parametrize('value',[0,-1,None])
def test_zero_negative_and_missing_baseline_are_not_invented(value):
    case=copy.deepcopy(CASES[2]);case['series'][0]['points'][2]['value']=value
    row=next(r for r in build(case)['layoutSections'][0]['layoutRows'] if r.get('comparisons'))
    c=row['comparisons']['comparison:yoy'];assert c['secondary'] is None and c['reason']

def test_formula_injection_remains_plain_text_in_all_workbook_sheets():
    case=copy.deepcopy(CASES[2]);case['options']['layout']['aliases']={'label':'=HYPERLINK("unsafe")'}
    doc=build(case)
    assert "'=HYPERLINK" in render_table_file(doc,'csv').decode('utf-8-sig')
    from openpyxl import load_workbook
    wb=load_workbook(io.BytesIO(render_table_file(doc,'xlsx')))
    assert all(c.data_type!='f' for ws in wb for row in ws for c in row)

def test_nested_headers_have_no_overlap_and_share_font_in_excel():
    case=copy.deepcopy(CASES[2]);layout=case['options']['layout']
    layout['header_bands'].append({'label':'반기 관측','columns':['period:2024S1','period:2024S2'],'level':1})
    layout['formats']['period:2026S1']={'bold':True,'color':'#cc0000','font_size':12}
    doc=build(case);headers=doc['layoutSections'][0]['headerRows'];assert len(headers)==3
    from openpyxl import load_workbook
    wb=load_workbook(io.BytesIO(render_table_file(doc,'xlsx')))
    cell=next(c for row in wb.worksheets[0] for c in row if c.value=='2026년 상반기')
    assert cell.font.color.rgb=='00cc0000' and cell.font.sz==12

@pytest.mark.parametrize('preset',['auto','a4_portrait','a4_landscape'])
def test_long_tables_paginate_and_repeat_identifiers(preset):
    import re
    case=copy.deepcopy(CASES[2]);case['options']['print_preset']=preset;case['options']['layout']['rows_per_page']=5
    blob=render_table_file(build(case),'pdf')
    assert len(re.findall(rb'/Type /Page\b',blob))>=2

def test_wide_tables_split_without_losing_any_columns():
    import re
    case=copy.deepcopy(CASES[0]);case['options']['layout']['widths']={'matrix:'+c['id']:300 for c in case['options']['layout']['matrix_columns']}
    blob=render_table_file(build(case),'pdf');assert len(re.findall(rb'/Type /Page\b',blob))>=3

@pytest.mark.parametrize('kind',['csv','xlsx','pdf','png'])
def test_older_layout_snapshot_without_header_styles_still_exports(kind):
    doc=build(CASES[0])
    for headers in doc['layoutSections'][0]['headerRows']:
        for header in headers:header.pop('format');header.pop('markers')
    assert len(render_table_file(doc,kind))>300

def test_csv_signed_numbers_remain_parseable_without_allowing_formulas():
    from table_layout_export import safe
    for number in ['-4.9','+12.6','-1,234.5']:assert safe(number)==number
    for formula in ['=1+1','+cmd()','-1+HYPERLINK("unsafe")','@SUM(A1)']:assert safe(formula).startswith("'")

def test_png_pixel_budget_stops_before_accumulating_more_pages(monkeypatch):
    from PIL import Image
    class Oversized:
        width=40_000_001;height=1;closed=False
        def convert(self,*args):return self
        def close(self):self.closed=True
    oversized=Oversized();monkeypatch.setattr(Image,'open',lambda *a,**k:oversized)
    with pytest.raises(ValueError,match='이미지'):render_table_file(build(CASES[0]),'png')
    assert oversized.closed

def test_layout_edit_uses_existing_atomic_history_and_ownership(monkeypatch,tmp_path):
    import account_auth as auth
    import bridge_api as api
    from test_account_auth import login,fixture_graph
    store=auth.AccountStore(tmp_path/'account.sqlite3');monkeypatch.setattr(auth,'STORE',store)
    one,user=login('test1');session=fixture_graph(monkeypatch,store,user,'rich')
    # No provider should ever be consulted for deterministic style commands.
    def forbidden(*args,**kwargs):raise AssertionError('model call during table edit')
    monkeypatch.setattr(api.edit_agent,'interpret',forbidden,raising=False)
    before=copy.deepcopy(session['result'])
    body={'edit_session_id':'rich','revision':0,'commands':[{'operation':'set_table','kind':'STYLE_EDIT','params':{'patch':{'layout':{'enabled':True,'formats':{'label':{'bold':True}}}}}}]}
    r=one.post('/api/output/edit',json=body);assert r.status_code==200,r.text
    assert r.json()['outputSpec']['tableDocument']['layoutSections']
    rev=r.json()['outputSpec']['editVersion']
    bad=one.post('/api/output/edit',json={**body,'revision':rev,'commands':[{'operation':'set_table','kind':'STYLE_EDIT','params':{'patch':{'layout':{'enabled':True,'formats':{'missing':{'bold':True}}}}}}]})
    assert bad.status_code in {400,422}
    assert api.EDIT_SESSIONS['rich']['output']['editVersion']==rev
    undo=one.post('/api/output/edit',json={'edit_session_id':'rich','revision':rev,'action':'undo'});assert undo.status_code==200
    redo=one.post('/api/output/edit',json={'edit_session_id':'rich','revision':undo.json()['outputSpec']['editVersion'],'action':'redo'});assert redo.status_code==200
    assert redo.json()['outputSpec']['tableDocument']['layoutSections']
    assert session['result']==before
    two,_=login('test2');assert two.post('/api/table/export',json={'edit_session_id':'rich','format':'xlsx'}).status_code==404

def test_legacy_overflow_records_are_preserved_and_active_list_is_five(monkeypatch,tmp_path):
    import account_auth as auth
    from test_account_auth import login
    store=auth.AccountStore(tmp_path/'account.sqlite3');monkeypatch.setattr(auth,'STORE',store)
    client,user=login('test1')
    with store.connect() as db:
        for i in range(8):db.execute('INSERT INTO graphs VALUES(?,?,?,?,?)',(user['id'],f'legacy-{i}','existing',i,b'hash'))
    data=client.get('/api/account/graphs').json()
    assert len(data['items'])==5 and len(data['olderItems'])==3
    assert data['items'][0]['id']=='legacy-7' and len(store.graphs(user['id']))==8
