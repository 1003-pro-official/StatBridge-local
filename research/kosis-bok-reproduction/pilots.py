"""Verify complete report contents against the supported local KOSIS tables.

Original PDFs provide comparison evidence only. All generated values come from
local KOSIS CSVs. Numerical/coverage/layout verdicts are deliberately separate.
"""
from __future__ import annotations
import csv
import json
import re
import sys
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
CACHE = ROOT / '.venv/cache/kosis-bok-reproduction'
sys.path[:0] = [str(ROOT / 'src/backend'), str(ROOT / 'src/agent')]
from statbridge_mcp.statistics_service import StatisticsService
from statbridge_mcp.data_snapshot import observation_version


def numbers(text):
    return [Decimal(x.replace(',', '')) for x in re.findall(r'[-+]?\d[\d,]*\.\d+', text or '')]


def complete(rows, periods):
    indexed = {r['PRD_DE']: Decimal(r['DT']) for r in rows}
    if len(indexed) != len(rows) or set(indexed) != set(periods) or any(not n.is_finite() for n in indexed.values()):
        raise ValueError('필수 관측치 누락·중복·비수치')
    return [indexed[p] for p in periods]


def main():
    import pdfplumber
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    import plotly.graph_objects as go
    service = StatisticsService()
    collection = json.loads((CACHE/'collection.json').read_text())
    cls = list(csv.DictReader((ROOT/'data/processed/bok_classifications.csv').open(encoding='utf-8-sig')))
    items = {r['TBL_ID']:r['ITM_ID'] for r in csv.DictReader((ROOT/'data/processed/bok_items.csv').open(encoding='utf-8-sig'))}
    out = CACHE/'original-pilots';out.mkdir(exist_ok=True)
    manifests=[];results=[]
    (out/'results.json').write_text(json.dumps(dict(status='running',started_at=datetime.now(timezone.utc).isoformat())))

    def source(ident, page):
        release = next(r for r in collection['releases'] if r['id']==ident)
        file = next(f for f in release['files'] if f['name'].endswith('.pdf'))
        return dict(release_id=ident, title=release['title'], url=release['url'], pdf_page=page, sha256=file['sha256']), ROOT/file['path']

    def series(table, code, label=None):
        candidates=[r for r in cls if r['TBL_ID']==table and (r['ITM_ID'].endswith('.'+code) if code else r['ITM_NM']==label)]
        if len(candidates)!=1:raise ValueError(f'계열 정의 모호: {table} {code} {label}')
        c=candidates[0]
        return dict(table_id=table,item_id=items[table],classifications={f"objL{c['OBJ_ID_SN']}":c['ITM_ID']},label=label or c['ITM_NM'])

    def data(target, periods, frequency):
        rows=service._filter_local_data(target['table_id'],target['item_id'],target['classifications'],frequency,min(periods),max(periods))
        rows=[r for r in rows if r['PRD_DE'] in periods]
        allowed_units = {'연%', '연리%'} if target['table_id'].startswith('DT_121') else {'십억원'} if target['table_id'].startswith('DT_151') else {'', '지수'}
        if any(r.get('UNIT_NM','') not in allowed_units for r in rows):
            raise ValueError('원본과 자료의 단위가 일치하지 않습니다.')
        target['observation_version']=observation_version(rows)
        return complete(rows, periods)

    def provenance_sheet(wb, src, definitions):
        ws=wb.create_sheet('자료 근거')
        for key,value in src.items():ws.append([key,str(value)])
        ws.append(['source','local_csv; KOSIS 한국은행 349개 표'])
        ws.append(['checked_at',datetime.now(timezone.utc).isoformat()])
        for target in definitions:
            ws.append([target['label'],target['table_id'],json.dumps(target['classifications'],ensure_ascii=False),target['observation_version']])
        ws.append(['verification','전체 수치 셀 대조; 픽셀 및 병합 헤더 동일성 미검증'])

    # All 13 rows, five nonconsecutive month columns and monthly changes.
    src,path=source('10095382',2)
    periods=['202312','202412','202509','202510','202511']
    defs=[series('DT_121Y002',c) for c in ['BEABAA2','BEABAA21','BEABAA22']]
    defs += [series('DT_121Y006',c) for c in ['BECBLA01','BECBLA02','BECBLA0201','BECBLA0202','BECBLA03','BECBLA0302']]
    defs += [series('DT_121Y013','BEABAB2'),series('DT_121Y015','BECBLB01')]
    values=[data(d,periods,'M') for d in defs]
    labels=[d['label'] for d in defs]
    values.insert(9,[values[3][i]-values[0][i] for i in range(5)]);labels.insert(9,'금리차(B－A)')
    values.append([values[11][i]-values[10][i] for i in range(5)]);labels.append('금리차(D－C)')
    matrix=[v+[v[-1]-v[-2]] for v in values]
    with pdfplumber.open(path) as doc:
        original_rows=[numbers(line) for line in doc.pages[1].extract_text().splitlines() if len(numbers(line))==6]
    if len(original_rows)!=13:raise ValueError('원본 금리표 전체 셀 추출 실패')
    original_columns=list(zip(*original_rows))
    comparisons=[abs(matrix[i][j]-original_columns[j][i])<=Decimal('0.005') for i in range(13) for j in range(6)]
    wb=Workbook();ws=wb.active;ws.title='예금은행 가중평균금리'
    ws.append(['예금은행 가중평균금리 (연 %, %p)']);ws.merge_cells('A1:G1')
    ws.append(['구분',*periods,'월중등락'])
    for label,row in zip(labels,matrix):ws.append([label,*map(float,row)])
    ws.append(['주: 2025.11 p는 잠정치. 원본 대비 전체 78개 수치 셀 대조.'])
    for row in ws.iter_rows(min_row=3,max_row=15,min_col=2):
        for cell in row:cell.number_format='0.00'
    for cell in ws[2]:cell.fill=PatternFill('solid',fgColor='DCE6F1');cell.font=Font(bold=True)
    ws.column_dimensions['A'].width=55
    for col in 'BCDEFG':ws.column_dimensions[col].width=16
    ws.freeze_panes='B3';ws.sheet_view.showGridLines=False
    provenance_sheet(wb,src,defs)
    wb.save(out/'RATE-FULL.xlsx')
    manifests.append(dict(id='RATE-FULL',source=src,periods=periods,series=defs,derived_rows=['B-A','D-C'],derived_column='202511-202510',unit='연 %, %p',notes=['202511 잠정치'],required_cells=78))
    results.append(dict(id='RATE-FULL',coverage_verified=True,numeric_cells=78,numeric_matches=sum(comparisons),numeric_verified=all(comparisons),presentation='행 순서·비연속 기간·등락·소수 2자리·잠정치 주석 반영; 원본 병합 헤더/픽셀 동일성 미검증',full_original_verified=False))

    # Complete original ESI figure range, both series, and benchmark 100.
    src,path=source('10095393',8)
    periods=[f'{y}{m:02}' for y in range(2013,2026) for m in range(1,13) if f'{y}{m:02}'<='202512']
    defs=[series('DT_513Y001',c) for c in ['E1000','E2000']]
    sample_periods=['202506','202507','202508','202509','202510','202511','202512']
    manifest=dict(comparison_periods=sample_periods,id='ESI-FULL',source=src,periods=periods,series=defs,benchmark=100,axes=dict(min=50,max=140,step=10),notes=['원본 기준치 100은 2003~2024년 중 장기평균; 현행 계열의 기준기간과 다를 수 있음','순환변동치는 계절 및 불규칙 변동요인 제거'],required_observations=312)
    try:
        values=[data(d,periods,'M') for d in defs]
        fig=go.Figure()
        for target,vals in zip(defs,values):fig.add_scatter(x=[p[:4]+'-'+p[4:] for p in periods],y=list(map(float,vals)),name=target['label'],mode='lines')
        fig.add_hline(y=100,line_dash='dot',annotation_text='기준치(100)')
        fig.update_layout(title='경제심리지수 추이 — KOSIS 현재 자료',yaxis=dict(range=[50,140],dtick=10),xaxis_title='월',template='plotly_white')
        fig.add_annotation(text='현행 KOSIS 자료: 2025년 공표 당시 수치와 불일치 · 과거 원본 재현 아님',xref='paper',yref='paper',x=0,y=-0.2,showarrow=False)
        fig.write_html(out/'ESI-CURRENT.html',include_plotlyjs=True)
        # The PDF contains only seven tabulated months; do not call the remaining image points numerically verified.
        with pdfplumber.open(path) as doc:text=doc.pages[7].extract_text()
        checks=[]
        for label,vals in zip(['ESI','ESI순환변동치'],values):
            line=next((line for line in text.splitlines() if re.sub(r'\s+','',line).startswith(label) and len(numbers(line))==7),None)
            if line:checks.extend(abs(a-b)<=Decimal('0.05') for a,b in zip([vals[periods.index(p)] for p in sample_periods],numbers(line)))
        results.append(dict(id='ESI-FULL',coverage_verified=True,observations=sum(map(len,values)),numeric_cells=len(checks),numeric_matches=sum(checks),numeric_verified=False,presentation='전체 기간·양 계열·기준선·축 범위 반영; 원본 이미지 전 시점 수치 및 픽셀 동일성 미검증',full_original_verified=False))
    except ValueError as exc:
        results.append(dict(id='ESI-FULL',coverage_verified=False,numeric_verified=False,full_original_verified=False,reason=str(exc)))
    manifests.append(manifest)

    # Household credit: all 20 columns, including two reference series.
    src,path=source('11063896',7)
    periods=['202304','202404','202504',*[f'{y}{q:02}' for y in [2024,2025] for q in range(1,5)],'202601','202602']
    defs=[series('DT_151Y001',c) for c in ['1000000','1100000']]
    defs += [series('DT_151Y004','11000A0'),series('DT_151Y001','1110000'),series('DT_151Y004','11100A0'),series('DT_151Y001','1111000'),series('DT_151Y001','11A0000'),series('DT_151Y001','1120000'),series('DT_151Y004','11200A0')]
    defs += [series('DT_151Y001',c) for c in ['1121000','1121400','1122000','1122300','1122400','1123000','1200000','1210000','1220000']]
    defs += [series('DT_151Y004',c) for c in ['1120093','1120094']]
    # Annual stock rows explicitly reuse Q4, never sum quarters.
    unique=list(dict.fromkeys(periods));values=[data(d,unique,'Q') for d in defs]
    matrix=[[values[j][unique.index(p)] for j in range(20)] for p in periods]
    with pdfplumber.open(path) as doc:text=doc.pages[6].extract_text()
    panels=[[numbers(line) for line in text.splitlines() if len(numbers(line))==n] for n in [9,11]]
    if any(len(panel)!=13 for panel in panels):raise ValueError('원본 가계신용 전체 셀 추출 실패')
    original=[a+b for a,b in zip(*panels)]
    # KOSIS stock unit is 십억원. No unit inference from label is permitted.
    comparisons=[abs(a-b)<=Decimal('0.05') for row,other in zip(matrix,original) for a,b in zip(row,other)]
    wb=Workbook();ws=wb.active;ws.title='가계신용 잔액'
    ws.append(['가계신용 잔액 (십억원)'])
    for start,end in [(0,9),(9,20)]:
        ws.append(['연분기말',*[d['label'] for d in defs[start:end]]])
        for i,(p,row) in enumerate(zip(periods,matrix)):
            ws.append([p[:4] if i<3 else p,*map(float,row[start:end])])
        ws.append([])
    for row in ws.iter_rows(min_row=3,min_col=2):
        for cell in row:
            if isinstance(cell.value,(int,float)):cell.number_format='#,##0.0'
    notes=['주: 2025~2026 잠정치. 연말 잔액은 4분기 말값.',
           '1) 생명보험회사, 손해보험회사, 우체국보험',
           '2) 공무원연금관리공단, 사립학교교직원연금공단 등',
           '3) 카드회사, 할부금융회사 등',
           '4) 주택도시기금, 한국주택금융공사 등',
           '5) 증권회사, 자산유동화회사(한국주택금융공사 유동화분 포함), 대부사업자(외감법인 등) 등',
           '6) 한국장학재단 등', '7) 개인대상 주택관련 정책대출']
    for note in notes:ws.append([note])
    ws.freeze_panes='B3';ws.sheet_view.showGridLines=False
    for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=22
    for ri in [2,17]:
        for cell in ws[ri]:
            cell.fill=PatternFill('solid',fgColor='DCE6F1');cell.font=Font(bold=True);cell.alignment=Alignment(wrap_text=True)
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth=1;ws.sheet_properties.pageSetUpPr.fitToPage=True
    provenance_sheet(wb,src,defs)
    wb.save(out/'CREDIT-FULL.xlsx')
    manifests.append(dict(id='CREDIT-FULL',source=src,periods=periods,series=defs,required_cells=260,unit='십억원',aggregation='연간 잔액=4분기 말값',notes=['보험·연금·여신전문·공적금융·기타중개·기타 기관의 원본 주석 1~7 유지 필요','2025~2026 잠정치']))
    results.append(dict(id='CREDIT-FULL',coverage_verified=True,numeric_cells=260,numeric_matches=sum(comparisons),numeric_verified=all(comparisons),presentation='20개 열·13개 시점·참고 2열 포함; 원본 2단 배치·기관 범위 주석 반영; 병합 헤더·픽셀 동일성 미검증',full_original_verified=False))
    (HERE/'원본재현명세.json').write_text(json.dumps(manifests,ensure_ascii=False,indent=2)+'\n')
    (out/'results.json').write_text(json.dumps(dict(status='complete',checked_at=datetime.now(timezone.utc).isoformat(),results=results),ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(results,ensure_ascii=False,indent=2))


if __name__=='__main__':main()
